from docx import Document

fp = r'生成结果\精修版\2023-2024-2《商务数据分析》教学设计 杜媛.docx'
doc = Document(fp)
t3 = doc.tables[3]

# 检查所有数据行的C3（能力目标列）是否被截断
print('=== C3 能力目标列 完整文本检查 ===')
for ri in range(3, len(t3.rows)):
    txt = t3.cell(ri, 3).text.strip()
    # 检查是否以不完整的句子结尾
    last_char = txt[-1] if txt else ''
    has_period = txt.endswith('。') or txt.endswith('：') or txt.endswith('）')
    print(f'  R{ri} [{len(txt):3d} chars] end="{txt[-10:]}" complete={has_period}')

print('\n=== C5 知识目标列 完整文本检查 ===')
for ri in range(3, len(t3.rows)):
    txt = t3.cell(ri, 5).text.strip()
    last_char = txt[-1] if txt else ''
    has_period = txt.endswith('。') or txt.endswith('：') or txt.endswith('）')
    print(f'  R{ri} [{len(txt):3d} chars] end="{txt[-10:]}" complete={has_period}')

print('\n=== C6 思政目标列 完整文本检查 ===')
for ri in range(3, len(t3.rows)):
    txt = t3.cell(ri, 6).text.strip()
    has_period = txt.endswith('。') or txt.endswith('：') or txt.endswith('）')
    print(f'  R{ri} [{len(txt):3d} chars] end="{txt[-10:]}" complete={has_period}')

print('\n=== C7 素质目标列 完整文本检查 ===')
for ri in range(3, len(t3.rows)):
    txt = t3.cell(ri, 7).text.strip()
    has_period = txt.endswith('。') or txt.endswith('：') or txt.endswith('）')
    print(f'  R{ri} [{len(txt):3d} chars] end="{txt[-10:]}" complete={has_period}')
