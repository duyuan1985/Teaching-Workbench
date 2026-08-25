"""调试实践内容写入问题"""
from docx import Document
from generate import write_cell, merge_v
import store

# 加载数据
offerings = store.rows('SELECT * FROM offerings WHERE id=20')
o = offerings[0]
tasks = store.rows('SELECT * FROM tasks WHERE offering_id=20 ORDER BY seq')
sessions = store.rows('SELECT * FROM sessions WHERE offering_id=20 ORDER BY week_no, lesson_date')
units = store.rows('SELECT * FROM curriculum_units WHERE offering_id=20 ORDER BY seq')

# 生成
from gen_plan import generate_plan
fp = generate_plan(o, tasks, sessions, units)

# 检查
doc = Document(fp)
t0 = doc.tables[0]

# 检查任务1的所有行
print("任务1区域（R2-R8）:")
for ri in range(2, 9):
    cells = t0.rows[ri].cells
    texts = [c.text.strip()[:20] for c in cells]
    print(f"  R{ri}: col3={texts[3][:20]} col4={texts[4]} col5={texts[5]} col6={texts[6]} col7={texts[7]} col8={texts[8]}")

# 用table.cell()检查
print("\n用table.cell()检查:")
for ri in range(2, 9):
    vals = []
    for ci in range(10):
        try:
            v = t0.cell(ri, ci).text.strip()[:15]
        except:
            v = "ERR"
        vals.append(v)
    print(f"  R{ri}: {vals}")
