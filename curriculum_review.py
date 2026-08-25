import re
from pathlib import Path

from docx import Document
from pptx import Presentation

from legacy_ppt import extract_slides

import store


EXCLUDED_PPT_TEXT = ("企业级卓越人才培养（信息类专业集群）", "企业级卓越人才培养")
NOISE_TEXT = (
    "任务技能", "天津滨海迅腾科技集团", "迅腾科技集团", "http://", "https://",
    "网页设计与制作—HTML5+CSS3项目实战", "网页设计与制作-HTML5+CSS3项目实战",
    "口令：", "___PPT", "谢谢！", "感谢聆听", "感谢观看", "ＴＨＡＮＫＳ",
    "目录", "学习目标", "学习路径", "任务描述", "系统环境",
    "LOGO", "logo", "Logo",
)

TITLE_CORRECTIONS = {
    "同城旅游界面设计": "同程旅游界面设计",
    "去哪儿旅游主界面设计": "携程旅游主界面设计",
    "酷狗音乐播放器界面": "酷狗音乐播放器界面设计",
    "开发迅腾国际集团首页": "HTML5+CSS3开发迅腾科技集团首页",
}

BUSINESS_PROJECT_TITLES = {
    1: "电子商务数据分析基础",
    2: "描述性统计分析",
    3: "市场数据分析",
    4: "竞争数据分析",
    5: "商品数据分析",
    6: "客户数据分析",
    7: "运营与销售数据分析",
    8: "库存数据分析",
}

PHOTOSHOP_SKILLS = {
    1: "Photoshop工作界面；文件新建、打开与存储；图像尺寸、分辨率与色彩模式；常用菜单与面板；图像文件格式与规范输出",
    2: "移动、复制、剪切与自由变换；画笔与橡皮擦工具；图像尺寸与画布调整；历史记录与基础修饰；基础编辑案例实施",
    3: "选区创建与调整；边界与羽化；亮度对比度、色阶与曲线；曝光度与色彩平衡；黑白效果与综合调色",
    4: "图层组织与混合模式；图层样式；通道创建、编辑、分离与合并；快速蒙版与图层蒙版；矢量蒙版与剪贴蒙版；非破坏性合成",
    5: "智能滤镜与滤镜库；Camera Raw与镜头校正；液化与消失点；模糊、风格化与综合特效；滤镜效果检查与规范输出",
}


def _clean_content(text):
    parts = [re.sub(r"\s+", "", part) for part in str(text or "").split("；") if part.strip()]
    parts = [
        part for part in parts
        if not any(excluded in part for excluded in EXCLUDED_PPT_TEXT)
        and not any(noise in part for noise in NOISE_TEXT)
        and not re.fullmatch(r"(?:PART\s*)?\d+", part, re.I)
    ]
    unique = []
    for part in parts:
        if part not in unique:
            unique.append(part)
    return "；".join(unique)


def _slides(text):
    result = []
    for match in re.finditer(r"第(\d+)页：(.*?)(?=第\d+页：|$)", str(text or ""), re.S):
        parts = [part.strip() for part in match.group(2).split("；") if part.strip()]
        result.append((int(match.group(1)), parts))
    return result


def _section(text, start_page, end_page):
    match = re.search(rf"第{start_page}页：(.*?)(?=第{end_page}页：|$)", text)
    return re.sub(r"\s+", " ", match.group(1)).strip() if match else ""


def _title(text, fallback):
    first = _section(text, 1, 2)
    parts = [part.strip() for part in first.split("；") if part.strip()]
    chapter = next((part for part in parts if re.search(r"第\s*\d+\s*章", part)), "")
    if chapter:
        title = re.sub(r"\s+", " ", chapter).strip()
        return TITLE_CORRECTIONS.get(title, title)
    parts = [part for part in parts if not re.fullmatch(r"项\s*目\s*[一二三四五六七八九十]+", part)]
    parts = [part for part in parts if not re.fullmatch(r"第?\s*(?:\d+|[一二三四五六七八九十]+)\s*章?", part)]
    parts = [part for part in parts if "企业级卓越人才培养" not in part]
    parts = [part for part in parts if not any(noise in part for noise in NOISE_TEXT)]
    title = parts[0] if parts else fallback
    return TITLE_CORRECTIONS.get(title, title)


def _objectives_and_skills(text):
    slides = _slides(text)
    objectives = ""
    for _, parts in slides:
        if parts and any("学习目标" in part for part in parts[:2]):
            raw_objectives = _clean_content("；".join(parts[1:]))
            objectives = re.sub(r"(?<!^)(?=(?:了解|掌握|熟悉|能够))", "；", raw_objectives)
            break
    topics = []
    for page, parts in slides:
        if page == 1 or not parts:
            continue
        topic = re.sub(r"_+PPT\d*", "", parts[0]).strip()
        if not topic or topic in ("学习目标", "谢谢！"):
            continue
        if any(noise in topic for noise in NOISE_TEXT):
            continue
        if len(topic) > 40 or topic in topics:
            continue
        topics.append(topic)
    skills = _clean_content("；".join(topics[:10]))
    return objectives, skills or objectives


def _concise_skills(text):
    result = []
    for part in str(text or "").split("；"):
        value = part.strip()
        if not value or len(value) > 32:
            continue
        if re.fullmatch(r"(?:任务|PART)\s*[\d一二三四五六七八九十-]+", value, re.I):
            continue
        if value.startswith("图") or any(noise in value for noise in NOISE_TEXT):
            continue
        if value not in result:
            result.append(value)
    return "；".join(result)


def _project_number(name):
    match = re.search(r"项目\s*0*(\d+)", name, re.I)
    if match:
        return int(match.group(1))
    chinese = re.search(r"项目\s*([一二三四五六七八九十]+)", name)
    values = {"一": 1, "二": 2, "三": 3, "四": 4, "五": 5, "六": 6, "七": 7, "八": 8, "九": 9, "十": 10}
    return values.get(chinese.group(1), 0) if chinese else 0


def _group_ppts(ppts, offering):
    course = offering["course_name"]
    textbook = f"{offering.get('textbook_version', '')} {offering.get('textbook_path', '')}"
    use_projects = course == "新媒体平台运营与推广" or "商务数据分析与决策" in textbook
    if not use_projects:
        return [[item] for item in ppts]
    groups = {}
    for item in ppts:
        parent = Path(item["file_path"]).parent
        number = _project_number(parent.name)
        key = (number, str(parent))
        groups.setdefault(key, []).append(item)
    return [groups[key] for key in sorted(groups, key=lambda value: (value[0] or 999, value[1]))]


def _group_title(items, offering, first_title):
    parent_name = Path(items[0]["file_path"]).parent.name
    number = _project_number(parent_name)
    textbook = f"{offering.get('textbook_version', '')} {offering.get('textbook_path', '')}"
    if offering["course_name"] == "新媒体平台运营与推广" and number:
        return re.sub(r"^项目\s*0*\d+\s*", "", parent_name).strip()
    if "商务数据分析与决策" in textbook and number:
        return BUSINESS_PROJECT_TITLES.get(number, f"项目{number}")
    if re.fullmatch(r"第\s*0*\d+\s*章", first_title):
        value = re.sub(r"^\d+\s*", "", parent_name).strip()
        return value or first_title
    return first_title


def _outline_projects(offering_id):
    documents = store.rows(
        "SELECT * FROM resource_items WHERE offering_id=? AND lower(file_path) LIKE ? ORDER BY file_path",
        (offering_id, "%.docx"),
    )
    candidates = [item for item in documents if "教学大纲" in item["title"] or "课程大纲" in item["title"]]
    for item in candidates:
        try:
            document = Document(item["file_path"])
        except Exception:
            continue
        projects = []
        for table in document.tables:
            for row in table.rows:
                text = "；".join(cell.text.strip() for cell in row.cells if cell.text.strip())
                match = re.search(r"项目\s*([一二三四五六七八九十]+|\d+)\s*([^（(；]+)[（(](\d+)学时[）)](.*)", text, re.S)
                if not match:
                    continue
                number = _project_number("项目" + match.group(1))
                title = re.sub(r"\s+", "", match.group(2)).strip()
                hours = int(match.group(3))
                tail = match.group(4)
                tasks = [
                    re.sub(r"^任务\s*[一二三四五六七八九十\d]+\s*", "", part).strip()
                    for part in re.split(r"[；\n]", tail)
                    if part.strip().startswith("任务")
                ]
                projects.append({
                    "number": number, "title": title, "hours": hours,
                    "skills": _clean_content("；".join(tasks)), "source_file": item["file_path"],
                })
        if projects:
            return sorted(projects, key=lambda value: value["number"])
    return []


def _modernization(title, objectives, skills, course_name=""):
    combined = f"{title} {objectives} {skills}"
    standards = "岗位质量规范、安全要求、数据与知识产权保护要求"
    technology = "教材配套工具与当前行业常用工具"
    process = "任务分析—方案设计—实施操作—检查测试—改进优化—成果交付"
    methods = "项目化教学、任务驱动、分层练习、合作学习、过程性评价"
    additions = []
    is_web = any(word in combined for word in ("HTML", "CSS", "网页", "网站", "Canvas"))
    if is_web:
        standards = "HTML语义化、W3C网页规范、网页无障碍、网络安全与数字版权要求"
        technology = "Visual Studio Code、浏览器开发者工具及移动端设备模拟"
        process = "需求分析—结构设计—编码实现—调试测试—迭代优化—成果交付"
    elif any(word in combined for word in ("Python", "函数", "模块", "类", "程序")):
        standards = "Python编码规范、软件安全、开源许可证与数据保护要求"
        technology = "Python 3、主流集成开发环境、虚拟环境、调试与测试工具"
        process = "问题分析—算法设计—编码实现—测试调试—重构优化—程序交付"
    elif any(word in combined for word in ("数据分析", "数据清洗", "可视化", "模型", "算法")):
        standards = "数据质量、数据安全、个人信息保护与分析结果可解释性要求"
        technology = "电子表格、Python数据分析工具、可视化与模型评估工具"
        process = "业务理解—数据获取—清洗处理—分析建模—结果解释—决策呈现"
    elif any(word in combined for word in ("新媒体", "运营", "推广", "营销", "内容")):
        standards = "平台运营规范、广告合规、个人信息保护与数字版权要求"
        technology = "内容生产、数据监测、平台运营与智能辅助工具"
        process = "用户分析—内容策划—生产发布—互动运营—数据复盘—迭代优化"
    elif course_name == "图形图像设计":
        standards = "数字图像规格、色彩管理、素材版权与作品输出规范"
        technology = "Adobe Photoshop、非破坏性编辑、图层与蒙版、智能对象和规范导出"
        process = "需求分析—素材整理—图层构建—图像处理—细节检查—格式输出—成果归档"
    elif course_name == "数据标注":
        standards = "数据标注规范、质量验收规则、个人信息保护与数据安全要求"
        technology = "图像、文本、语音和视频标注工具及质量复核方法"
        process = "任务解读—样本预处理—规范标注—质量复核—问题纠正—数据交付"
    if any(word in combined for word in ("固定布局", "流动布局", "导航", "布局")):
        additions.append("使用Flex/Grid和媒体查询完成响应式布局")
        technology += "、Flex、Grid和媒体查询"
    if is_web and any(word in combined for word in ("图像", "背景", "图片")):
        additions.append("补充响应式图片、替代文本、素材压缩与版权规范")
    if "表单" in combined or "注册" in combined:
        additions.append("补充HTML5表单验证、个人信息最小化收集和安全提示")
        standards += "、个人信息保护"
    if any(word in combined for word in ("audio", "video", "音频", "视频")):
        additions.append("补充多媒体格式兼容、字幕/替代内容与播放控制")
    if "Canvas" in combined or "钟表" in combined:
        additions.append("结合JavaScript模块化组织、设备像素比和性能调试")
    if any(word in combined for word in ("网站", "首页", "集团")):
        additions.append("采用移动优先、组件化思维和跨视口验收")
    revised = _clean_content("；".join([skills] + additions) if skills else "；".join(additions))
    rationale = "保留教材项目载体和核心知识，替换落后开发环境，补充当前岗位使用的标准、工具、流程与质量要求。"
    return revised, rationale, standards, technology, process, methods


def _allocate_hours(weights, total_hours):
    if not weights:
        return []
    units = total_hours // 2
    total_weight = sum(weights)
    raw = [units * weight / total_weight for weight in weights]
    # One teaching session is two periods. Dense textbooks may therefore use
    # two-hour chapter units; forcing four hours makes the blueprint exceed
    # the official course total when a course has many short chapters.
    base = [max(1, int(value)) for value in raw]
    while sum(base) > units:
        index = max((i for i, value in enumerate(base) if value > 1), key=lambda i: base[i] - raw[i], default=None)
        if index is None:
            break
        base[index] -= 1
    while sum(base) < units:
        index = max(range(len(raw)), key=lambda i: raw[i] - base[i])
        base[index] += 1
    return [value * 2 for value in base]


def _slide_weight(path):
    path = Path(path)
    try:
        if path.suffix.lower() == ".ppt":
            return max(1, len(extract_slides(path)))
        if path.suffix.lower() == ".pptx":
            return max(1, len(Presentation(path).slides))
    except Exception:
        pass
    return 1


def build_curriculum_review(offering):
    """构建课程蓝本审查，支持AI模式和配置规则模式。"""
    ai_mode = store.get_setting("ai_curriculum_review", "0") == "1"
    if ai_mode:
        from ai_curriculum_review import ai_review_curriculum
        return ai_review_curriculum(offering["id"])

    rule = _get_review_rule(offering["course_name"])
    if rule:
        return _review_with_rule(offering, rule)

    return _review_default(offering)


def _get_review_rule(course_name):
    """查询课程的蓝本审查配置规则"""
    rows = store.rows(
        "SELECT * FROM curriculum_review_rules WHERE course_name=? AND is_active=1",
        (course_name,),
    )
    if not rows:
        return None
    return rows[0]


def _review_with_rule(offering, rule):
    """使用配置规则进行蓝本审查"""
    ppts = store.rows(
        "SELECT * FROM resource_items WHERE offering_id=? "
        "AND resource_type IN ('PPT课件','旧版PPT课件') ORDER BY file_path",
        (offering["id"],),
    )
    if not ppts:
        raise ValueError("资源索引中没有可读取的PPT课件。")

    main_ppts = [
        item for item in ppts
        if re.search(r"[\\/]0?\d+\s*课程\s*ppt[\\/]", item["file_path"], re.I)
    ]
    if main_ppts:
        ppts = main_ppts

    group_mode = rule["ppt_group_mode"]
    project_pattern = rule["project_pattern"] or r"项目\s*0*(\d+)"
    if group_mode == "project_dir":
        groups = {}
        for item in ppts:
            parent = Path(item["file_path"]).parent
            match = re.search(project_pattern, parent.name, re.I)
            number = int(match.group(1)) if match else 0
            key = (number, str(parent))
            groups.setdefault(key, []).append(item)
        groups = [groups[key] for key in sorted(groups, key=lambda v: (v[0] or 999, v[1]))]
    elif group_mode == "merge_all":
        groups = [ppts]
    else:
        groups = [[item] for item in ppts]

    outline_keyword = rule["outline_keyword"] or "教学大纲"
    outline = _outline_projects(offering["id"]) if outline_keyword else []
    outline_by_number = {item["number"]: item for item in outline}
    if outline:
        ppt_by_number = {}
        for group in groups:
            number = _project_number(Path(group[0]["file_path"]).parent.name)
            if number:
                ppt_by_number[number] = group
        groups = [
            ppt_by_number.get(item["number"], [{
                "file_path": item["source_file"], "content_excerpt": "", "title": item["title"],
            }])
            for item in outline
        ]

    units = []
    weights = []
    specified_hours = []

    objective_keyword = rule["objective_keyword"] or "学习目标"
    skill_keywords = [s.strip() for s in (rule["skill_keywords"] or "").split(",") if s.strip()]

    for index, items in enumerate(groups, 1):
        item = items[0]
        excerpt = item["content_excerpt"]
        first_title = _title(excerpt, item["title"])
        group_number = _project_number(Path(item["file_path"]).parent.name)
        outline_item = outline_by_number.get(group_number) or (outline[index - 1] if outline and index <= len(outline) else None)
        title = outline_item["title"] if outline_item else _group_title(items, offering, first_title)

        objectives_parts, skills_parts = [], []
        if outline_item and outline_item["skills"]:
            skills_parts.append(outline_item["skills"])
        for grouped_item in items:
            grouped_excerpt = grouped_item["content_excerpt"]
            grouped_objectives, grouped_skills = _objectives_and_skills_custom(grouped_excerpt, objective_keyword, skill_keywords)
            if grouped_objectives:
                objectives_parts.append(grouped_objectives)
            if grouped_skills:
                skills_parts.append(grouped_skills)

        objectives = _clean_content("；".join(objectives_parts))
        skills = _clean_content("；".join(skills_parts))

        revised, rationale, standards, technology, process, methods = _modernization(title, objectives, skills, offering["course_name"])
        skill_count = min(5, max(1, len([part for part in skills.split("；") if part.strip()])))
        weight = skill_count
        weights.append(weight)
        specified_hours.append(outline_item["hours"] if outline_item else None)
        source_file = str(Path(item["file_path"]).parent) if len(items) > 1 else item["file_path"]
        units.append((title, source_file, objectives, skills, revised, rationale, standards, technology, process, methods))

    assessment_hours = min(int(offering.get("weekly_hours") or 4), int(offering["total_hours"]))
    if specified_hours and all(value is not None for value in specified_hours) and sum(specified_hours) + assessment_hours == int(offering["total_hours"]):
        hours = specified_hours
    else:
        hours = _allocate_hours(weights, int(offering["total_hours"]) - assessment_hours)

    units.append((
        "综合评价与课程总结", "", "课程成果提交、综合评价与学习总结",
        "课程成果汇报与评价；课程总结与复习",
        "课程成果汇报与评价；课程总结与复习", "完成课程成果验收、问题复盘与后续学习规划。",
        "课程质量标准与成果评价规范", "成果展示、文档整理与评价记录工具",
        "成果整理—展示汇报—多元评价—复盘改进", "成果答辩、同伴互评、自我评价和总结反思",
    ))
    hours.append(assessment_hours)

    with store.connect() as db:
        db.execute("DELETE FROM curriculum_units WHERE offering_id=?", (offering["id"],))
        for index, (unit, suggested_hours) in enumerate(zip(units, hours), 1):
            db.execute(
                """INSERT INTO curriculum_units
                (offering_id,seq,project_title,source_file,source_objectives,source_skills,review_action,
                 revised_focus,rationale,new_standards,new_technology,new_process,new_methods,suggested_hours,approval_status)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (offering["id"], index, unit[0], unit[1], unit[2], unit[3], "规则审查",
                 *unit[4:], suggested_hours, "待确认"),
            )
        db.commit()
    return len(units), hours


def _objectives_and_skills_custom(text, objective_keyword, skill_keywords):
    """使用自定义关键字提取学习目标和技能点"""
    if not text:
        return "", ""
    slides = _slides(text)
    if not slides:
        return "", ""

    objectives = []
    skills = []

    for slide in slides[:5]:
        if objective_keyword in slide:
            start = slides.index(slide)
            for next_slide in slides[start:start + 3]:
                parts = [p.strip() for p in re.split(r"[；;。\n]", next_slide) if p.strip()]
                for part in parts:
                    if objective_keyword not in part and not any(n in part for n in NOISE_TEXT):
                        objectives.append(part)

    for slide in slides:
        for kw in skill_keywords:
            if kw in slide:
                parts = [p.strip() for p in re.split(r"[；;。\n]", slide) if p.strip()]
                for part in parts:
                    if kw not in part and not any(n in part for n in NOISE_TEXT):
                        skills.append(part)

    return _clean_content("；".join(objectives)), _clean_content("；".join(skills))


def _review_default(offering):
    """原有蓝本审查逻辑"""
    ppts = store.rows(
        "SELECT * FROM resource_items WHERE offering_id=? AND resource_type IN ('PPT课件','旧版PPT课件') ORDER BY file_path",
        (offering["id"],),
    )
    if not ppts:
        raise ValueError("资源索引中没有可读取的PPT课件。")
    main_ppts = [
        item for item in ppts
        if re.search(r"[\\/]0?\d+\s*课程\s*ppt[\\/]", item["file_path"], re.I)
    ]
    if main_ppts:
        ppts = main_ppts
    groups = _group_ppts(ppts, offering)
    outline = _outline_projects(offering["id"])
    outline_by_number = {item["number"]: item for item in outline}
    if outline:
        ppt_by_number = {}
        for group in groups:
            number = _project_number(Path(group[0]["file_path"]).parent.name)
            if number:
                ppt_by_number[number] = group
        groups = [
            ppt_by_number.get(item["number"], [{
                "file_path": item["source_file"], "content_excerpt": "", "title": item["title"],
            }])
            for item in outline
        ]
    units = []
    weights = []
    specified_hours = []
    for index, items in enumerate(groups, 1):
        item = items[0]
        excerpt = item["content_excerpt"]
        first_title = _title(excerpt, item["title"])
        group_number = _project_number(Path(item["file_path"]).parent.name)
        outline_item = outline_by_number.get(group_number) or (outline[index - 1] if outline and index <= len(outline) else None)
        title = outline_item["title"] if outline_item else _group_title(items, offering, first_title)
        objectives_parts, skills_parts = [], []
        if outline_item and outline_item["skills"]:
            skills_parts.append(outline_item["skills"])
        for grouped_item in items:
            grouped_excerpt = grouped_item["content_excerpt"]
            is_chapter_textbook = "网页设计与制作—HTML5+CSS3项目实战" in grouped_item["file_path"]
            if is_chapter_textbook:
                grouped_objectives = _clean_content(_section(grouped_excerpt, 2, 3))
                grouped_skills = grouped_objectives
            else:
                grouped_objectives, grouped_skills = _objectives_and_skills(grouped_excerpt)
            if grouped_objectives:
                objectives_parts.append(grouped_objectives)
            skill_title = re.sub(r"^技能点\s*[一二三四五六七八九十\d]+\s*", "", _title(grouped_excerpt, grouped_item["title"])).strip()
            concise_title = _concise_skills(skill_title)
            if len(items) > 1 and concise_title:
                skills_parts.append(concise_title)
            if grouped_skills and offering["course_name"] != "新媒体平台运营与推广":
                concise = _concise_skills(grouped_skills) if len(items) > 1 else grouped_skills
                if concise:
                    skills_parts.append(concise)
        objectives = _clean_content("；".join(objectives_parts))
        skills = _clean_content("；".join(skills_parts))
        if offering["course_name"] == "图形图像设计" and index in PHOTOSHOP_SKILLS:
            skills = PHOTOSHOP_SKILLS[index]
        revised, rationale, standards, technology, process, methods = _modernization(title, objectives, skills, offering["course_name"])
        skill_count = min(5, max(1, len([part for part in skills.split("；") if part.strip() and "企业级卓越" not in part])))
        weight = skill_count + (2 if any(word in title for word in ("钟表", "集团首页")) else 0)
        if offering["course_name"] == "Python程序设计":
            weight = 1
        elif offering["course_name"] == "图形图像设计":
            weight = sum(_slide_weight(grouped_item["file_path"]) for grouped_item in items)
        weights.append(weight)
        specified_hours.append(outline_item["hours"] if outline_item else None)
        source_file = str(Path(item["file_path"]).parent) if len(items) > 1 else item["file_path"]
        units.append((title, source_file, objectives, skills, revised, rationale, standards, technology, process, methods))
    assessment_hours = min(int(offering.get("weekly_hours") or 4), int(offering["total_hours"]))
    if specified_hours and all(value is not None for value in specified_hours) and sum(specified_hours) + assessment_hours == int(offering["total_hours"]):
        hours = specified_hours
    else:
        hours = _allocate_hours(weights, int(offering["total_hours"]) - assessment_hours)
    units.append((
        "综合评价与课程总结", "", "课程成果提交、综合评价与学习总结",
        "课程成果汇报与评价；课程总结与复习",
        "课程成果汇报与评价；课程总结与复习", "完成课程成果验收、问题复盘与后续学习规划。",
        "课程质量标准与成果评价规范", "成果展示、文档整理与评价记录工具",
        "成果整理—展示汇报—多元评价—复盘改进", "成果答辩、同伴互评、自我评价和总结反思",
    ))
    hours.append(assessment_hours)
    with store.connect() as db:
        db.execute("DELETE FROM curriculum_units WHERE offering_id=?", (offering["id"],))
        for index, (unit, suggested_hours) in enumerate(zip(units, hours), 1):
            db.execute(
                """INSERT INTO curriculum_units
                (offering_id,seq,project_title,source_file,source_objectives,source_skills,review_action,
                 revised_focus,rationale,new_standards,new_technology,new_process,new_methods,suggested_hours,approval_status)
                VALUES (?,?,?,?,?,?,?,?,?,?,?,?,?,?,?)""",
                (offering["id"], index, unit[0], unit[1], unit[2], unit[3], "更新", *unit[4:], suggested_hours, "待确认"),
            )
        db.commit()
    return len(units), hours
